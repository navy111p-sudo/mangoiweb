# -*- coding: utf-8 -*-
"""마이마이·강사 대상 영문 메일 2건 (.docx)

① 마이마이에게 — 8/13 배포분 확인 요청 3가지
② 강사 전체에게 — 사전 연기 지급 변경 공지 (8/15 시행)

본문은 영어(받는 사람이 필리핀 사람), 사장님용 메모만 한국어 회색으로 붙인다.
그대로 복사해서 메일에 붙일 수 있도록 «메일 본문» 과 «메모» 를 확실히 구분한다.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()
st = DOC.styles['Normal']
st.font.name = 'Malgun Gothic'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

for s in DOC.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)


def shade(par, hexcolor):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    pPr.append(shd)


def p(text='', size=10.5, bold=False, color=None, before=0, after=5,
      italic=False, align=None, fill=None, indent=0, font=None):
    par = DOC.add_paragraph()
    if align: par.alignment = align
    pf = par.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.3
    if indent: pf.left_indent = Cm(indent)
    r = par.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.name = font or 'Malgun Gothic'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font or 'Malgun Gothic')
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if fill: shade(par, fill)
    return par


def rich(parts, size=10.5, before=0, after=5, indent=0, fill=None):
    """[(텍스트, bold), ...] 를 한 문단에 — 문장 중간 강조용."""
    par = DOC.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.3
    if indent: pf.left_indent = Cm(indent)
    for txt, bold in parts:
        r = par.add_run(txt)
        r.font.size = Pt(size); r.bold = bold
        r.font.name = 'Malgun Gothic'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    if fill: shade(par, fill)
    return par


def memo(text):
    """사장님용 한국어 메모 — 메일 본문이 아님을 색으로 구분."""
    p('📌 ' + text, size=9, color='808080', italic=True, after=8)


def h_mail(no, to, subject):
    p('', size=4, after=2)
    p('MAIL %d — %s' % (no, to), size=15, bold=True, color='1F3864', before=4, after=3)
    p('Subject:  ' + subject, size=11, bold=True, color='C00000', after=2, fill='FFF2CC')
    p('─' * 74, size=8, color='BFBFBF', after=8)


def rate_table(rows):
    t = DOC.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hd = t.rows[0].cells
    for i, txt in enumerate(('Situation', 'Pay')):
        par = hd[i].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(txt); r.bold = True; r.font.size = Pt(10)
        r.font.name = 'Malgun Gothic'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
        shade(par, 'D9E2F3')
    for left, right, strong in rows:
        c = t.add_row().cells
        pl = c[0].paragraphs[0]
        rl = pl.add_run(left); rl.font.size = Pt(10)
        rl.font.name = 'Malgun Gothic'
        rl._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
        pr = c[1].paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pr.add_run(right); rr.font.size = Pt(10); rr.bold = strong
        rr.font.name = 'Malgun Gothic'
        rr._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
        if strong: shade(pr, 'FCE4EC')
    for row in t.rows:
        row.cells[0].width = Cm(11.0)
        row.cells[1].width = Cm(4.5)
    return t


# ═══════════════════════════ 표지 ═══════════════════════════
p('Mangoi — Emails to Send', size=19, bold=True, color='1F3864',
  align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p('보낼 메일 2건 — 영문 본문 (그대로 복사해서 쓰세요)', size=10,
  color='808080', align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p('August 13, 2026', size=10, color='808080',
  align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

p('이 파일에는 두 통이 들어 있습니다.', size=10.5, bold=True, after=4)
rich([('① 마이마이에게 ', True), ('— 오늘 배포한 수정 3가지를 확인해 달라는 요청.', False)], after=3, indent=0.4)
rich([('② 강사 전체에게 ', True), ('— 사전 연기 지급 변경 공지 (8월 15일 토요일 시행).', False)], after=8, indent=0.4)
p('회색 이탤릭 줄(📌)은 사장님께 드리는 메모입니다 — 메일에 넣지 마세요.',
  size=9.5, color='808080', after=6, fill='F2F2F2')

DOC.add_page_break()

# ═══════════════════════ ① 마이마이 확인 요청 ═══════════════════════
h_mail(1, 'To: Maimai', "You were right about the duplicated pages — fixed, please check")

p('Hi Maimai,', after=8)
p('Thank you for going through all the questions — your answers were genuinely useful. '
  'Two of them changed what we did, so I want to tell you exactly what happened.', after=8)

rich([('First, an apology — twice over. ', True),
      ('When you first told us the textbooks were still slow, the fixes were not actually live yet, '
       'so you were testing the old version. And then it turned out we had been chasing the wrong '
       'problem entirely: ', False),
      ('the books were never slow. They were duplicated.', True),
      (' You told us exactly that, with exact numbers, and you were right. Thank you for pushing '
       'back instead of letting it go.', False)], after=8)

rich([('Everything below is now live.', True),
      (' Could you please check these three things when you get a chance?', False)], after=10)

p('1.  Do the books now have the correct number of pages?', size=11.5, bold=True, color='1F3864', after=3)
rich([('You were right, and we were wrong. ', True),
      ('We had assumed the books were slow to load, and we spent yesterday making them load faster. '
       'Then you wrote ', False),
      ('"THE BOOK IS NOT SLOW ONLY THAT THE PAGES AT THE LIBRARY TRIPLED OR DOUBLED"', True),
      (' — and that was the real problem all along.', False)], after=4, indent=0.5)
rich([('We checked the database and your numbers were exact. ', False),
      ('BTS 1 001 had 115 pages in the library, but only 23 of them were real', True),
      (' — the rest were the same pages uploaded over and over. Across the whole library, '
       '38,922 pages were really only 17,170. Some books were duplicated 5 times.', False)],
     after=4, indent=0.5)
rich([('We have now removed the duplicates. ', True),
      ('BTS 1 001 is 23 pages. BTS 2 went from 762 pages down to 246. ', False),
      ('Nothing was deleted — every file is still stored, and we can put any of it back instantly. '
       'We also fixed the uploader so that uploading the same book twice can no longer double it.', False)],
     after=4, indent=0.5)
rich([('👉 Please open a few books and tell me whether the page counts look right now. '
       'If any book still looks doubled, just tell me ', False), ('the book name', True),
      (' — I can check it directly.', False)],
     size=10, after=9, indent=0.5, fill='FFF9E6')

p('2.  Are the MES books gone from the library?', size=11.5, bold=True, color='1F3864', after=3)
rich([('You confirmed ', False), ('"Level 1 to 7 are MES"', True),
      (', so we hid exactly those. Nothing was deleted — the files and all past lesson records are '
       'untouched, and we can bring any of them back instantly if you find one you still need.', False)],
     after=9, indent=0.5)

p('3.  Does the page list stay open, and can you resize it?', size=11.5, bold=True, color='1F3864', after=3)
p('Picking a page no longer closes the window — only the [✕] does. You can drag the bottom-right '
  'corner to make it bigger, and it will remember the size next time.', after=10, indent=0.5)

p('One more thing — about the level tests.', size=11.5, bold=True, color='1F3864', after=3)
rich([('You wrote ', False),
      ('"Teachers can upload every level test but we don\'t know how to add in the library."', True),
      (' You were not doing anything wrong. ', False), ('The screen was lying to you.', True),
      (' It said "✅ Saved!" even when the file only went to that teacher\'s own computer and never '
       'reached the shared library. Teachers had no way to know.', False)], after=5)
rich([('It now tells the truth: if a file cannot go into the shared library, you will see ', False),
      ('"⚠️ Saved to THIS COMPUTER only"', True),
      (' with an explanation of what to do next. For now, please send level-test files to the head '
       'office to upload, and we will let you know if we open that up to teachers directly.', False)], after=10)

p('Thank you again for reporting these so clearly — the detail in your notes is what made it '
  'possible to find the real causes.', after=8)
p('Best regards,', after=2)
p('[Your name]', after=10)

memo('「사과」 문단이 부담스러우면 빼셔도 됩니다. 다만 «느리다고 했는데 왜 또 묻나» 싶을 수 있어 '
     '배포가 안 된 상태였다는 걸 먼저 밝히는 편이 신뢰에 낫습니다.')
memo('사전 연기 0페소 이야기는 이 메일에 일부러 넣지 않았습니다 — 아래 ② 공지로 따로 보내세요.')

DOC.add_page_break()

# ═══════════════════════ ② 강사 공지 ═══════════════════════
h_mail(2, 'To: All Teachers  (cc. Maimai)',
        'Notice — change to postponed class pay, starting Saturday, August 15')

p('Dear Teachers,', after=8)
rich([('Starting ', False), ('Saturday, August 15, 2026', True),
      (', there is a change to how postponed classes are paid. Please read this carefully, '
       'as it affects your pay.', False)], after=10)

p('The rule from August 15', size=11.5, bold=True, color='1F3864', after=5)
p('Pay for a postponed class will depend on when the class was postponed:', after=6)
rate_table([
    ('Class completed', '₱50', False),
    ('Postponed WITHIN 30 minutes of the start time', '₱50  (full pay)', False),
    ('Postponed MORE THAN 30 minutes before the start time', '₱0', True),
])
p('', size=6, after=6)

p('Why', size=11.5, bold=True, color='1F3864', after=3)
p('When a class is postponed well in advance, we still have time to fill that slot with another '
  'class. When it is postponed at the last minute, that time cannot be used — so you are still '
  'paid in full in that case.', after=9)

p('What this means for you', size=11.5, bold=True, color='1F3864', after=3)
p('Nothing changes for classes you actually teach, or for last-minute postponements. This only '
  'affects classes postponed with plenty of notice.', after=5)
p('Examples — for a class scheduled at 7:00 PM:', size=10, italic=True, color='595959', after=3)
rich([('•  Postponed at ', False), ('6:45 PM', True), (' (15 minutes before)  →  ', False),
      ('₱50, full pay', True)], size=10, after=2, indent=0.6)
rich([('•  Postponed at ', False), ('6:20 PM', True), (' (40 minutes before)  →  ', False),
      ('₱0', True)], size=10, after=2, indent=0.6)
rich([('•  Postponed at ', False), ('2:00 PM', True), (' (5 hours before)  →  ', False),
      ('₱0', True)], size=10, after=10, indent=0.6)

p('Good to know', size=11.5, bold=True, color='1F3864', after=3)
rich([('•  ', False), ('Classes up to Friday, August 14 are not affected.', True),
      (' The new rule starts with classes on Saturday, August 15.', False)], size=10, after=3, indent=0.4)
rich([('•  ', False), ('Your class fee page shows the reason for every amount.', True),
      (' If a class shows ₱0, you can see whether it was recorded as an early postponement. '
       'If something looks wrong, tell your manager — we can check the records.', False)],
     size=10, after=3, indent=0.4)
rich([('•  ', False), ('A class postponed by the student or by the office is not counted against you.', True),
      (' If you see otherwise, please raise it with your manager.', False)], size=10, after=3, indent=0.4)
rich([('•  ', False), ('Nothing else changes', True),
      (' — the existing rules for late arrivals and for feedback stay the same.', False)],
     size=10, after=10, indent=0.4)

p('Please pass this on to all teachers today so everyone has notice before Saturday. If anyone has '
  'questions, they can speak with Maimai or reply to this message. We would rather answer questions '
  'now than have anyone surprised on payday.', after=8)
p('Thank you for your work.', after=8)
p('[Your name]', after=1)
p('Mangoi', after=10)

memo('운영 한 가지 — 규칙이 8/13 배포로 이미 켜져 있습니다. 공지는 8/15 시행이므로, '
     '관리자 → 급여 → 공제 규칙에서 [사전 연기 지급률]을 지금 꺼 두었다가 15일 아침에 다시 켜면 '
     '목·금 이틀치가 어긋나지 않습니다. (테스트 사이트라 실제 급여 영향이 없다면 안 하셔도 무방합니다.)')
memo('8월 15일은 토요일이고 오늘이 목요일이라, 강사들에게 이틀 여유가 있습니다.')

OUT = '/home/user/mangoiweb/docs/마이마이_메일_2건_2026-08-13.docx'
DOC.save(OUT)
print('saved:', OUT)
