# -*- coding: utf-8 -*-
"""마이마이에게 보낼 질문지 (.docx)

강사·매니저가 필리핀 사람이라 영어를 본문으로 하고, 한국어를 회색 보조줄로 함께 둔다.
(사장님이 같은 파일 하나로 내용을 확인하실 수 있어야 한다)
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

# ── 기본 서식: 한글이 들어가므로 동아시아 폰트까지 지정해야 네모(두부)가 안 뜬다 ──
style = DOC.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

for s in DOC.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)


def shade(par, hexcolor):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    pPr.append(shd)


def para(text, size=10.5, bold=False, color=None, space_before=0, space_after=4,
         italic=False, align=None, fill=None, indent=0):
    p = DOC.add_paragraph()
    if align: p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    if indent: pf.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.name = 'Malgun Gothic'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if fill: shade(p, fill)
    return p


def qblock(no, tag, en_title, ko_title, en_body, ko_body, answer_hint=None):
    """질문 한 덩어리 — 영어가 본문, 한국어가 보조."""
    para('', size=4, space_after=2)
    para('Q%d.  %s' % (no, en_title), size=12, bold=True, color='1F3864',
         space_before=6, space_after=2, fill='EAF0F8')
    para('[%s]  %s' % (tag, ko_title), size=9.5, bold=True, color='7F7F7F', space_after=5)
    para(en_body, size=10.5, space_after=3)
    para(ko_body, size=9.5, color='808080', space_after=4)
    if answer_hint:
        para('✎  %s' % answer_hint, size=10, bold=True, color='C00000',
             space_after=3, fill='FFF2CC')
    # 답 적는 칸
    para('Your answer: ' + '_' * 74, size=10, color='999999', space_after=2)
    para('_' * 88, size=10, color='999999', space_after=10)


# ═══════════════════════ 표지 ═══════════════════════
para('Mangoi — Questions for Maimai', size=19, bold=True, color='1F3864',
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('NEW TEACHER\'S PAGE PROBLEM — follow-up questions', size=11.5,
     color='C00000', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, bold=True)
para('마이마이에게 물어볼 것 — 강사페이지 지적사항 후속', size=10,
     color='808080', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
para('August 13, 2026', size=10, color='808080',
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

para('Hi Maimai — thank you for the detailed reports. We fixed 25 of your 31 items. '
     'Below are the only things we could not finish on our own, because they need '
     'your answer or your decision. Please reply next to each question.',
     size=10.5, space_after=3, fill='F2F2F2')
para('마이마이가 보낸 31건 중 25건을 고쳤습니다. 아래는 «마이마이의 답이 있어야» '
     '끝낼 수 있는 것들입니다.', size=9.5, color='808080', space_after=14)

# ═══════════════════════ 질문 ═══════════════════════
qblock(
    1, '가장 급함 · 8/13 ①',
    'Which textbooks should we hide from the library?',
    '라이브러리에서 어느 교재를 숨길까요?',
    'You said "There are still MES books at the library, server textbooks". We found why our '
    'earlier fix did not work: it only hid books whose NAME contains "MES" — but none of the '
    'server textbooks have "MES" in their name. So nothing was hidden.\n\n'
    'We did NOT want to guess, because if we guess wrong, a book you actually use would '
    'disappear for every teacher. So we built a screen instead: '
    'Textbook Uploader → "🙈 Hide from library" → tick the ones to hide. '
    'It applies immediately and you can untick it any time. Files are never deleted.\n\n'
    'These are the server textbook groups. Which ones are the old MES books?\n'
    '   LEVEL 1 · LEVEL 3 · LEVEL 4 · LEVEL 5 · LEVEL 6 · LEVEL 7\n'
    '   Mangoi Books · Shake it up 2-18-25\n'
    '   004. I visited my grandparents · 007. Commercials · (기타 / Others)',
    '8월 11일 숨김이 안 통한 이유 — 필터가 «이름에 MES 가 든 것»만 걸렀는데, 서버 교재 '
    '이름에는 MES 가 하나도 없었습니다. 코드가 「LEVEL 1~7 이 MES 겠지」 하고 찍으면 틀렸을 때 '
    '강사가 쓰는 교재가 사라지므로, 사람이 고르는 화면을 만들었습니다.',
    'Please tick them yourself in the uploader screen — or just list the names here.')

qblock(
    2, '8/10 ① · 재현 안 됨',
    'Can you send a screen recording of the flickering?',
    '깜빡이는 화면을 영상으로 찍어 보내 주실 수 있나요?',
    'You reported "The classroom has flickering display" (Aug 10) and later "Only the top '
    'portion is flickering" (Aug 11).\n\n'
    'The Aug 11 one we found and fixed — the top chips were blinking every 1.2 seconds, and '
    'it only happened in ENGLISH mode, which is why our Korean staff never saw it.\n\n'
    'But we could not reproduce the Aug 10 one ("the classroom" in general). '
    'A 10-second screen recording would let us find it the same way.',
    '8/11 「위쪽만」은 찾아서 고쳤습니다(영어 모드에서만 나던 문제였습니다). '
    '8/10 「교실 전체」는 아직 재현을 못 했습니다.',
    'Is it still happening after Aug 12? If yes → a short screen recording, please.')

qblock(
    3, '8/7 ⑩',
    'Do you still need a participant list, or is the mic bar on each face enough?',
    '참가자 목록이 따로 필요할까요, 아니면 얼굴마다 붙은 음량 막대로 충분할까요?',
    'You wrote: "I don\'t know if it was mine that is working or the student\'s mic … '
    '[I want] participants to see if the student and teachers mic are working."\n\n'
    'We removed the floating mic meter (it was covering the screen-split panel) and now the '
    'volume bar sits INSIDE each person\'s video tile — so the bar under your face is your mic, '
    'and the bar under the student\'s face is theirs.\n\n'
    'Is that clear enough now? Or do you still want a separate list of names with mic status, '
    'like the old classroom had?',
    '떠 있던 미터를 없애고 얼굴 타일 안쪽으로 옮겼습니다. 이걸로 구분이 되는지, '
    '아니면 옛 교실처럼 «이름 + 마이크 상태» 목록이 따로 필요한지 물어봅니다.',
    'Tile bar is enough  /  I still want a participant list  — please circle one.')

qblock(
    4, '8/7-2 ① · 램 구매 전',
    'Are the textbooks faster now? (before we buy more RAM)',
    '교재가 지금은 빨라졌나요? (램을 사기 전에 확인)',
    'You reported the library was slow five separate times, and suggested buying 16GB RAM.\n\n'
    'We checked the actual database: 38,860 of 38,922 textbook files (99.8%) are ALREADY JPG '
    '— only 62 are PDF. So the problem was never "PDF is slow". The real cause is that each '
    'page is a separate file, so every ▶ downloaded a new file.\n\n'
    'Today we made it load 3 pages ahead and 1 page back, instead of just 1 ahead.\n\n'
    'Please use it for a few days and tell us: is it noticeably faster? '
    'If it is, you may not need to buy RAM at all.',
    '느림의 원인이 PDF 가 아니라 «1장=1파일이라 넘길 때마다 새로 받는 것» 이었습니다. '
    '미리 받는 양을 4배로 늘렸습니다. 램 구매 전에 며칠 써 보고 알려 달라고 요청합니다.',
    'After a few days: still slow  /  much better  — and on which book?')

qblock(
    5, '8/12 ④',
    'Who will upload the BTS placement test into the library?',
    'BTS 배치고사(레벨테스트) 교재는 누가 올릴까요?',
    'You asked: "Can we include BTS placement test (Level test) in the library?" and sent a '
    'Google Drive link.\n\n'
    'There is no placement test in the library at all right now. This is not a code change — '
    'the file just needs to be uploaded through the Textbook Uploader, the same way as any '
    'other book.\n\n'
    'Would you like to upload it yourself, or should our head office do it? '
    'Also — what name should the group have, so teachers can find it easily?',
    '코드 문제가 아니라 «파일을 올리는 일» 입니다. 누가 올릴지와 교재 묶음 이름만 정하면 됩니다.',
    'I will upload it  /  Please upload it for me  — and the name should be: ______')

qblock(
    6, '8/7 ④ · 확인',
    'Is the postpone setup OK as it is now?',
    '연기(postpone) 구조가 지금 상태로 괜찮은지 확인 부탁드립니다.',
    'You asked to remove the postpone option from the teacher\'s page, so that "managers and '
    'students can only postpone".\n\n'
    'We removed the [Postpone] button that sat on each lesson row — that was the one teachers '
    'could hit by accident.\n\n'
    'But the "Request a schedule change" card is still there, on purpose (owner\'s decision). '
    'A teacher can only REQUEST it — nothing changes until a manager approves. '
    'Does that match what you wanted, or should teachers not be able to request at all?',
    '실수로 누를 위험이 있던 [연기] 버튼은 없앴고, «요청»카드는 사장님 지시로 남겼습니다'
    '(승인해야 확정). 이 구조가 마이마이 뜻과 맞는지 확인합니다.',
    'This is fine  /  Teachers should not even be able to request  — please circle one.')

qblock(
    7, '8/7-2 ④ · 급하지 않음',
    'Where exactly did you see the MES sample at the class entrance?',
    '입장 화면의 MES 샘플 교재를 정확히 어디서 보셨나요?',
    'You noted "The sample on the page was MES not BTS" (and wrote "No changes" yourself, so '
    'we know this is not urgent).\n\n'
    'We checked every image in the code and there is no MES sample picture — the entrance '
    'screen shows the Mr. Mango character with a greeting in Korean, English and Chinese. '
    'So what you saw is probably uploaded textbook data, not a fixed image.\n\n'
    'If you see it again, a screenshot would help us find it.',
    '코드에는 MES 샘플 이미지가 없습니다. 올려 둔 교재 데이터 쪽으로 보입니다. '
    '원문에도 "No changes" 라고 적혀 있어 급하지 않습니다.',
    'Only if you see it again — a screenshot, please.')

# ═══════════════════════ 확인 요청 목록 ═══════════════════════
DOC.add_page_break()
para('Please also confirm these are working', size=15, bold=True, color='1F3864', space_after=2)
para('아래는 «고쳤습니다» 라고 보고된 것들입니다 — 실제로 되는지 확인 부탁드립니다',
     size=9.5, color='808080', space_after=8)
para('We fixed these but could not verify them ourselves, because the teacher and admin '
     'screens require a login. Please tick the ones that work, and tell us any that still do not.',
     size=10.5, space_after=3)
para('강사·관리자 화면은 로그인해야 보여서 저희가 눈으로 확인하지 못했습니다.',
     size=9.5, color='808080', space_after=8)

CHECKS = [
    ('Aug 13', 'The page list stays open when you pick a page (only [✕] closes it)',
     '페이지를 골라도 목록 창이 안 닫힙니다'),
    ('Aug 13', 'You can drag the bottom-right corner of the page list to make it bigger',
     '페이지 목록 창 오른쪽 아래 모서리를 끌면 커집니다'),
    ('Aug 13', 'Books load faster when you press ▶ (we now preload 3 pages ahead)',
     '▶ 를 눌렀을 때 더 빨리 넘어갑니다'),
    ('Aug 13', 'The video lessons page no longer opens on the MES tab',
     '동영상 강의를 열면 MES 탭이 아니라 BTS 로 시작합니다'),
    ('Aug 12', 'BTS books are now in order 1, 2, 3 … 34 (not 1, 10, 11, 2)',
     'BTS 교재가 1~34 순서대로 나옵니다'),
    ('Aug 12', 'Your camera stays ON when you open another browser tab',
     '다른 탭을 열어도 강사 카메라가 안 꺼집니다'),
    ('Aug 12', 'You can read the student\'s chat messages',
     '학생 채팅 글씨가 보입니다'),
    ('Aug 11', 'You can remove a wrongly-picked book with [✕] and [🗑] in the page list',
     '잘못 고른 교재를 페이지 목록에서 뺄 수 있습니다'),
    ('Aug 11', 'MES books and BTS/SIU/Teacher videos no longer appear during class',
     '수업 중 MES 교재와 BTS·SIU·강사 영상이 안 보입니다'),
    ('Aug 7', 'Level test classes now show in "My classes" with a LEVEL TEST badge',
     '레벨테스트가 「내 수업」에 배지와 함께 나옵니다'),
    ('Aug 7', 'You can enter a class early, and re-enter after it ends (all day)',
     '수업 전 미리 입장 · 끝난 뒤 다시 입장이 됩니다'),
    ('Aug 7', 'The weekly schedule has a date box to jump to any year/month/day',
     '주간 스케줄에서 날짜를 골라 바로 이동할 수 있습니다'),
    ('Aug 7', 'The mic indicator no longer covers the screen-split settings',
     '마이크 표시가 화면 분할 설정을 안 가립니다'),
    ('Aug 7', 'When the faces are hidden, a big [🙂 Show faces] button appears',
     '얼굴을 감추면 되살리는 큰 버튼이 뜹니다'),
    ('Aug 7', 'The class fee table shows time, student, status and deductions',
     '수업료 표에 시간·학생·상태·공제가 나옵니다'),
    ('Aug 7', 'Break time: you can set start/end yourself, and delete it by clicking the card',
     '휴식시간을 직접 정하고, 카드를 눌러 지울 수 있습니다'),
    ('Admin', 'Admin pages no longer jump to the middle when you pick a menu',
     '관리자 메뉴를 골라도 가운데로 안 밀립니다'),
]

tbl = DOC.add_table(rows=1, cols=3)
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for c, t in ((0, 'OK?'), (1, 'When'), (2, 'What should work now')):
    p = hdr[c].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(9.5)
    r.font.name = 'Malgun Gothic'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    shade(p, 'D9E2F3')

for when, en, ko in CHECKS:
    cells = tbl.add_row().cells
    p0 = cells[0].paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run('☐'); r0.font.size = Pt(13)
    p1 = cells[1].paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(when); r1.font.size = Pt(8.5); r1.font.color.rgb = RGBColor.from_string('808080')
    p2 = cells[2].paragraphs[0]
    r2 = p2.add_run(en); r2.font.size = Pt(9.5)
    r2.font.name = 'Malgun Gothic'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    r3 = p2.add_run('\n' + ko); r3.font.size = Pt(8.5)
    r3.font.color.rgb = RGBColor.from_string('A6A6A6')
    r3.font.name = 'Malgun Gothic'
    r3._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

# 열 너비
for row in tbl.rows:
    row.cells[0].width = Cm(1.3)
    row.cells[1].width = Cm(1.9)
    row.cells[2].width = Cm(13.0)

para('', size=8, space_after=6)
para('If any line above does NOT work, please write the book name / class / time so we can '
     'find it in the logs. Thank you! 🙏', size=10.5, bold=True, space_before=6, space_after=3)
para('안 되는 항목이 있으면 교재 이름·수업·시각을 적어 주세요 — 기록에서 찾겠습니다.',
     size=9.5, color='808080')

OUT = '/home/user/mangoiweb/docs/마이마이_질문지_2026-08-13.docx'
DOC.save(OUT)
print('saved:', OUT)
