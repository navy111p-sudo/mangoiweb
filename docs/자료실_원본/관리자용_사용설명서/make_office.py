# 관리자용 사용설명서 — PPTX·XLSX 만들기 (build.mjs 가 부릅니다)
#   python3 -I make_office.py <content.json> <캡처폴더> <출력폴더>
# content.json 은 build.mjs 가 content.mjs 에서 뽑아 넘깁니다(글 정본은 content.mjs 하나).
import json, sys, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.drawing.image import Image as XImage

content = json.load(open(sys.argv[1], encoding='utf-8'))
CAP, OUT = sys.argv[2], sys.argv[3]
os.makedirs(OUT, exist_ok=True)

def strip(s):
    import re
    s = re.sub(r'<[^>]+>', '', s)
    return s.replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')

for lang, doc in content.items():
    # ---------- PPTX ----------
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    ORANGE = RGBColor(0xF5, 0x9E, 0x0B); DARK = RGBColor(0x11, 0x18, 0x27); GRAY = RGBColor(0x4B, 0x55, 0x63)
    def tb(sl, x, y, w, h, text, size, color=DARK, bold=False):
        box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        return tf
    s = prs.slides.add_slide(blank)
    tb(s, 0.8, 2.3, 11.7, 1.2, strip(doc['title']), 40, DARK, True)
    tb(s, 0.8, 3.6, 11.7, 0.8, strip(doc['sub']), 20, GRAY)
    tb(s, 0.8, 4.4, 11.7, 0.6, strip(doc['date']), 16, ORANGE, True)
    s = prs.slides.add_slide(blank)
    tb(s, 0.6, 0.4, 12, 0.8, strip(doc['tocTitle']), 30, DARK, True)
    tf = tb(s, 0.8, 1.3, 11.8, 5.8, '', 16)
    for i, ch in enumerate(doc['chapters']):
        p = tf.add_paragraph() if i else tf.paragraphs[0]
        p.text = f"{i+1}. {strip(ch['title'])}"; p.font.size = Pt(16)
    s = prs.slides.add_slide(blank)
    tb(s, 0.6, 0.4, 12, 0.8, strip(doc['beforeTitle']), 28, DARK, True)
    tf = tb(s, 0.8, 1.3, 11.8, 5.8, '', 16)
    for i, x in enumerate(doc['before']):
        p = tf.add_paragraph() if i else tf.paragraphs[0]; p.text = '• ' + strip(x); p.font.size = Pt(16)
    for ci, ch in enumerate(doc['chapters']):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(7.5)); bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0xFF, 0xF7, 0xED); bar.line.fill.background()
        tb(s, 0.8, 2.6, 11.7, 1.2, f"{ci+1}. {strip(ch['title'])}", 36, DARK, True)
        tb(s, 0.8, 3.8, 11.7, 1.4, strip(ch.get('intro', '')), 18, GRAY)
        for st in ch['steps']:
            s = prs.slides.add_slide(blank)
            tb(s, 0.4, 0.25, 12.5, 0.7, f"{ci+1}. {strip(ch['title'])} — {strip(st['h'])}", 22, DARK, True)
            img = os.path.join(CAP, f"{lang}-{st['img']}.jpg") if st.get('img') else None
            if img and os.path.exists(img):
                from PIL import Image
                w, h = Image.open(img).size
                maxw, maxh = 7.9, 6.3
                sc = min(maxw / (w / 96), maxh / (h / 96))
                iw, ih = w / 96 * sc, h / 96 * sc
                s.shapes.add_picture(img, Inches(0.35), Inches(1.05), Inches(iw), Inches(ih))
                tx = 0.35 + iw + 0.25
            else:
                tx = 0.5
            tf = tb(s, tx, 1.05, 13.0 - tx, 6.2, '', 14)
            for i, x in enumerate(st['p']):
                p = tf.add_paragraph() if i else tf.paragraphs[0]; p.text = '• ' + strip(x); p.font.size = Pt(13 if len(st['p']) > 5 else 14); p.space_after = Pt(6)
    prs.save(os.path.join(OUT, doc['files']['pptx']))
    # ---------- XLSX ----------
    wb = Workbook(); ws = wb.active; ws.title = doc['xlsx']['sheet1']
    hdr = doc['xlsx']['header']
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='F59E0B')
    for ci, ch in enumerate(doc['chapters']):
        for si, st in enumerate(ch['steps']):
            ws.append([f"{ci+1}", strip(ch['title']), f"{ci+1}-{si+1}", strip(st['h']), '\n'.join('• ' + strip(x) for x in st['p']), st.get('img', '') and f"{doc['xlsx']['picSheet']}!{ci+1}-{si+1}"])
    for col, w in zip('ABCDEF', [6, 26, 7, 38, 100, 16]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws2 = wb.create_sheet(doc['xlsx']['picSheet'])
    r = 1
    for ci, ch in enumerate(doc['chapters']):
        for si, st in enumerate(ch['steps']):
            img = os.path.join(CAP, f"{lang}-{st['img']}.jpg") if st.get('img') else None
            if not (img and os.path.exists(img)): continue
            ws2.cell(row=r, column=1, value=f"{ci+1}-{si+1}  {strip(st['h'])}").font = Font(bold=True, size=13)
            from PIL import Image
            w, h = Image.open(img).size
            xi = XImage(img); sc = min(900 / w, 1); xi.width = int(w * sc); xi.height = int(h * sc)
            ws2.add_image(xi, f"A{r+1}")
            r += int(xi.height / 20) + 4
    ws2.column_dimensions['A'].width = 120
    ws3 = wb.create_sheet(doc['xlsx']['faqSheet'])
    ws3.append(doc['xlsx']['faqHeader'])
    for c in ws3[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='F59E0B')
    for q, a in doc['faq']: ws3.append([strip(q), strip(a)])
    ws3.column_dimensions['A'].width = 50; ws3.column_dimensions['B'].width = 110
    for row in ws3.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(os.path.join(OUT, doc['files']['xlsx']))
    print('office', lang)
