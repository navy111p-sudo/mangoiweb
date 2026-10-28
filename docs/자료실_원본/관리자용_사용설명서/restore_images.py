# 📘 관리자용 사용설명서 — 이미 배포된 PPTX 에서 단계 그림을 되살린다 (2026-10-28)
#   왜 — 장 하나만 고칠 때 39장을 전부 다시 찍으면, 그 사이 바뀐 화면이 섞여 «글과 그림이 다른» 설명서가 된다.
#        그래서 고치지 않는 장의 그림은 지금 자료실 PPTX 에 들어 있는 원본 JPG 를 그대로 꺼내 쓴다.
#   쓰는 법(리포 루트):
#     node -e "import('./docs/자료실_원본/관리자용_사용설명서/content.mjs').then(m=>process.stdout.write(JSON.stringify(m.T)))" > <c.json>
#     python3 -I "docs/자료실_원본/관리자용_사용설명서/restore_images.py" <c.json> cloudflare-deploy/public/library/admin <캡처폴더>
#   대조는 슬라이드 제목(«장번호. 장 제목 — 단계 제목»)으로 한다 — 순서로 맞추면 단계 하나만 늘어도 전부 한 칸씩 밀린다.
#   ⚠️ 제목을 바꾼 단계·새 단계는 PPTX 에 없으니 «못 찾음» 으로 나온다 → 그 그림만 capture.mjs 로 찍는다.
import json, sys, os
from pptx import Presentation
content = json.load(open(sys.argv[1], encoding='utf-8'))
LIB, OUT = sys.argv[2], sys.argv[3]
import re
strip = lambda s: re.sub(r'<[^>]+>', '', s).replace('&lt;','<').replace('&gt;','>').replace('&amp;','&')
for lang, fn in (('ko', 'admin-kr.pptx'), ('en', 'admin-en.pptx')):
    doc = content[lang]
    want = {}
    for ci, ch in enumerate(doc['chapters']):
        for st in ch['steps']:
            want[f"{ci+1}. {strip(ch['title'])} — {strip(st['h'])}"] = st['img']
    prs = Presentation(os.path.join(LIB, fn)); got = {}; unmatched = []
    for sl in prs.slides:
        pics = [s for s in sl.shapes if s.shape_type == 13]
        if not pics: continue
        title = next((s.text_frame.text for s in sl.shapes if s.has_text_frame and s.text_frame.text.strip()), '')
        img = want.get(title)
        if not img: unmatched.append(title); continue
        blob = pics[0].image.blob; ext = pics[0].image.ext
        assert ext in ('jpg', 'jpeg'), ext
        p = os.path.join(OUT, f'{lang}-{img}.jpg')
        if img in got and got[img] != blob: print('다른 그림이 같은 이름', lang, img)
        got[img] = blob; open(p, 'wb').write(blob)
    need = set(want.values())
    print(lang, '되살림', len(got), '/ 필요', len(need), '· 못 찾음(새로 찍을 것)', sorted(need - set(got)), '· 대조 안 된 슬라이드', unmatched)
