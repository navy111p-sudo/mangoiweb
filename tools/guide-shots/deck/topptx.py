#!/usr/bin/env python3
# 슬라이드 그림(01.jpg …)을 파워포인트 한 벌로 묶는다 — 자료실에서 내려받는 .pptx.
#   쓰기: python3 tools/guide-shots/deck/topptx.py <그림폴더> <내보낼 pptx> <제목>
#   한 장에 그림 한 장을 꽉 채운다(16:9). 편집용이 아니라 «보고 인쇄하는» 용도다.
import sys, glob, os
from pptx import Presentation
from pptx.util import Emu

src, out = sys.argv[1], sys.argv[2]
title = sys.argv[3] if len(sys.argv) > 3 else 'Mangoi'
files = sorted(glob.glob(os.path.join(src, '[0-9][0-9].jpg')))
if not files:
    sys.exit('그림을 찾지 못했습니다: ' + src)

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333in × 7.5in (16:9)
blank = prs.slide_layouts[6]
for f in files:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.core_properties.title = title
prs.save(out)
print(f'✅ PPTX {len(files)}장 → {out} ({os.path.getsize(out)//1024} KB)')
