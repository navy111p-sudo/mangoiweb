# -*- coding: utf-8 -*-
"""대상별 사용설명서 Word(docx) 생성 — 표지·소개·섹션(캡처)·표·차트데이터·FAQ·체크리스트·마무리"""
import os, importlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image as PILImage

C = importlib.import_module(os.environ.get("MANUAL_MODULE", "content"))
HERE = os.path.dirname(os.path.abspath(__file__))
OPT = os.path.join(HERE, "assets", "opt")
CHAR = os.path.join(HERE, "assets", "mango_char.png")
FONT = getattr(C, "FONT", "Malgun Gothic")   # 중국어판(content_zh)은 'Microsoft YaHei' 를 쓴다

def rgb(hexs): return RGBColor.from_string(hexs.lstrip("#")[:6])

def set_cell_bg(cell, hexs):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexs.lstrip("#")[:6])
    cell._tc.get_or_add_tcPr().append(sh)

def no_borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "none"); borders.append(e)
    tblPr.append(borders)

def md_runs(para, text, size=10.5, color="#334155", bold_color=None):
    """**bold** 파싱해 run 추가"""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg: continue
        r = para.add_run(seg); r.font.size = Pt(size); r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        bold = (i % 2 == 1)
        r.font.bold = bold
        r.font.color.rgb = rgb(bold_color if (bold and bold_color) else color)

def para(doc, text="", size=10.5, color="#334155", bold=False, align=None, space_after=4, space_before=0):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(space_before)
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); r.font.color.rgb = rgb(color)
    return p

def callout(doc, label, text, fill, label_color, text_color="#334155"):
    """색 배경 박스(1셀 표)"""
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tb.cell(0, 0); set_cell_bg(cell, fill); no_borders(tb)
    cell.width = Inches(6.4)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    lr = p.add_run(label + " "); lr.font.bold = True; lr.font.size = Pt(10.5); lr.font.name = FONT
    lr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); lr.font.color.rgb = rgb(label_color)
    md_runs(p, text, size=10, color=text_color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_image(doc, key, width=6.2):
    p = os.path.join(OPT, key + ".jpg")
    if not os.path.exists(p): return
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(4); pic.paragraph_format.space_after = Pt(8)
    run = pic.add_run(); run.add_picture(p, width=Inches(width))

def heading(doc, text, size, color, space_before=10, space_after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = True; r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); r.font.color.rgb = rgb(color)
    return p

def build(aud, out_path):
    t = C.THEMES[aud]; m = C.MANUALS[aud]
    p_, p2, acc = t["primary"], t["primary2"], t["accent"]
    title, subtitle = C.COVERS[aud]
    doc = Document()
    # 기본 폰트
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7); sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)

    # ── 표지 ──
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    para(doc, f"{t['emoji']} {C.BRAND['name']} · MANGOI", 13, p2, bold=True)
    para(doc, title, 30, p_, bold=True, space_after=6)
    para(doc, subtitle, 14, "#475569", bold=True, space_after=10)
    if os.path.exists(CHAR):
        pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pic.add_run().add_picture(CHAR, width=Inches(2.4))
    para(doc, f"{C.UI['target']} · {t['who']}", 11, "#334155", space_before=8)
    para(doc, f"🌐 {C.SITE}    📞 1644-0561    ✉ support@mangoi.co.kr", 10, "#64748B")
    doc.add_page_break()

    # ── 소개 ──
    intro = m["intro"]
    heading(doc, f"📘 {intro['title']}", 18, p_)
    md_runs(para(doc, "", 11), intro["body"], size=11, color="#334155")
    for ic, ttl, d in intro["points"]:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(3)
        r = pp.add_run(f"{ic} {ttl} — "); r.font.bold = True; r.font.size = Pt(11); r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); r.font.color.rgb = rgb(p_)
        r2 = pp.add_run(d); r2.font.size = Pt(10.5); r2.font.name = FONT
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); r2.font.color.rgb = rgb("#334155")
    doc.add_page_break()

    # ── 섹션 ──
    heading(doc, f"🧭 {C.UI['section']}", 16, p_, space_after=8)
    for s in m["sections"]:
        heading(doc, f"{s['no']:02d}.  {s['icon']} {s['title']}", 14, p_, space_before=12, space_after=2)
        para(doc, s["sub"], 10.5, "#64748B", space_after=6)
        if s.get("desc"):
            callout(doc, "📖", s["desc"], t["soft"], p_)
        para(doc, C.UI["howto"], 11, p_, bold=True, space_after=2)
        for i, step in enumerate(s["steps"], 1):
            pp = doc.add_paragraph(); pp.paragraph_format.left_indent = Inches(0.2); pp.paragraph_format.space_after = Pt(2)
            nr = pp.add_run(f"{i}. "); nr.font.bold = True; nr.font.size = Pt(10.5); nr.font.name = FONT
            nr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); nr.font.color.rgb = rgb(p2)
            md_runs(pp, step, size=10.5, color="#334155", bold_color=p_)
        if s.get("points"):
            para(doc, C.UI["point"], 10.5, p_, bold=True, space_before=4, space_after=2)
            for pt in s["points"]:
                pp = doc.add_paragraph(style=None); pp.paragraph_format.left_indent = Inches(0.2); pp.paragraph_format.space_after = Pt(1)
                br = pp.add_run("• "); br.font.color.rgb = rgb(p_); br.font.name = FONT
                md_runs(pp, pt, size=10, color="#334155", bold_color=p_)
        if s.get("note"):
            callout(doc, C.UI["note"], s["note"], "#FFFBEB", "#B45309", "#92670A")
        if s.get("tip"):
            callout(doc, C.UI["tip"], s["tip"], "#FFF7E6", C.BRAND["mango_dark"], "#5B4A2A")
        if s.get("imgs"):
            for k, _ in s["imgs"]:
                add_image(doc, k, width=3.0 if len(s["imgs"]) > 1 else 6.2)
        elif s.get("img"):
            add_image(doc, s["img"])

    doc.add_page_break()

    # ── 기능 요약 표 ──
    ft = m["feature_table"]
    heading(doc, f"📋 {ft['title']}", 16, p_)
    table = doc.add_table(rows=1, cols=len(ft["head"])); table.style = "Table Grid"
    for j, hd in enumerate(ft["head"]):
        cell = table.rows[0].cells[j]; set_cell_bg(cell, p_)
        rp = cell.paragraphs[0].add_run(hd); rp.font.bold = True; rp.font.size = Pt(10); rp.font.color.rgb = rgb("#FFFFFF")
        rp.font.name = FONT; rp._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for i, row in enumerate(ft["rows"]):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            if i % 2: set_cell_bg(cells[j], "#F5F7FB")
            rp = cells[j].paragraphs[0].add_run(str(val)); rp.font.size = Pt(9.5); rp.font.name = FONT
            rp._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); rp.font.color.rgb = rgb("#334155")
            if j == 0: rp.font.bold = True

    # ── 데이터(차트) : 데이터 표 ──
    heading(doc, C.UI["data"], 16, p_, space_before=14)
    for ch in list(m["charts"].values()):
        para(doc, ch["title"], 11, "#1E2340", bold=True, space_before=6, space_after=2)
        mx = max(ch["values"]) or 1
        tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
        for lab, val, col in zip(ch["labels"], ch["values"], (ch["colors"]*len(ch["values"]))[:len(ch["values"])]):
            cells = tb.add_row().cells
            r0 = cells[0].paragraphs[0].add_run(str(lab)); r0.font.size = Pt(9.5); r0.font.name = FONT
            r0._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); r0.font.color.rgb = rgb("#334155")
            # 값 셀: 막대 느낌(값에 비례한 블록 문자) + 숫자
            bar = "█" * max(1, int(round(val / mx * 18)))
            pc = cells[1].paragraphs[0]
            rb = pc.add_run(bar + " "); rb.font.size = Pt(9); rb.font.color.rgb = rgb(col); rb.font.name = FONT
            rn = pc.add_run(str(val)); rn.font.size = Pt(9.5); rn.font.bold = True; rn.font.color.rgb = rgb(col); rn.font.name = FONT
        tb.columns[0].width = Inches(2.0); tb.columns[1].width = Inches(4.2)

    doc.add_page_break()

    # ── FAQ ──
    heading(doc, C.UI["faq"], 16, p_)
    for q, a in m["faq"]:
        pq = doc.add_paragraph(); pq.paragraph_format.space_after = Pt(1)
        rq = pq.add_run(f"Q. {q}"); rq.font.bold = True; rq.font.size = Pt(11); rq.font.color.rgb = rgb(p_); rq.font.name = FONT
        rq._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pa = doc.add_paragraph(); pa.paragraph_format.space_after = Pt(6); pa.paragraph_format.left_indent = Inches(0.15)
        ra = pa.add_run(f"A. {a}"); ra.font.size = Pt(10.5); ra.font.color.rgb = rgb("#475569"); ra.font.name = FONT
        ra._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ── 체크리스트 ──
    cl = m["checklist"]
    heading(doc, f"✅ {cl['title']}", 16, p_, space_before=12)
    for it in cl["items"]:
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(2)
        rb = pp.add_run("☐  "); rb.font.size = Pt(12); rb.font.color.rgb = rgb(p_); rb.font.name = FONT
        md_runs(pp, it, size=11, color="#334155")

    # ── 마무리(결론) ──
    if m.get("conclusion"):
        cc = m["conclusion"]
        doc.add_page_break()
        heading(doc, f"🏁 {cc['title']}", 17, p_)
        md_runs(para(doc, "", 11), cc["lead"], size=11, color="#334155")
        heading(doc, C.UI["concl_flow"], 13, p_, space_before=8, space_after=4)
        tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
        for stg, act in cc["flow"]:
            cells = tb.add_row().cells; set_cell_bg(cells[0], t["soft"])
            r0 = cells[0].paragraphs[0].add_run(stg); r0.font.bold = True; r0.font.size = Pt(10); r0.font.color.rgb = rgb(p_); r0.font.name = FONT
            r0._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            md_runs(cells[1].paragraphs[0], act, size=10, color="#334155", bold_color=p_)
        tb.columns[0].width = Inches(1.7); tb.columns[1].width = Inches(4.5)
        heading(doc, C.UI["concl_principles"], 13, p_, space_before=8, space_after=4)
        for i, x in enumerate(cc["principles"], 1):
            pp = doc.add_paragraph(); pp.paragraph_format.left_indent = Inches(0.2); pp.paragraph_format.space_after = Pt(2)
            nr = pp.add_run(f"{i}. "); nr.font.bold = True; nr.font.size = Pt(10.5); nr.font.color.rgb = rgb(p2); nr.font.name = FONT
            md_runs(pp, x, size=10.5, color="#334155", bold_color=p_)
        heading(doc, C.UI["concl_remember"], 13, p_, space_before=8, space_after=4)
        for x in cc["remember"]:
            pp = doc.add_paragraph(); pp.paragraph_format.left_indent = Inches(0.2); pp.paragraph_format.space_after = Pt(2)
            br = pp.add_run("• "); br.font.color.rgb = rgb(acc); br.font.name = FONT
            md_runs(pp, x, size=10, color="#334155", bold_color=p_)
        callout(doc, "🍊", cc["closing"], t["soft"], p_)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print("DOCX", aud, "->", out_path)

if __name__ == "__main__":
    suffix = getattr(C, "OUT_SUFFIX", "")
    outdir = os.path.join(HERE, "build_docx" + suffix)
    for aud in list(C.MANUALS.keys()):
        build(aud, os.path.join(outdir, C.FILENAMES[aud] + ".docx"))
