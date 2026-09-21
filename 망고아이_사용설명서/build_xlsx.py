# -*- coding: utf-8 -*-
"""대상별 사용설명서 Excel 생성 (색인·단계별·체크리스트·FAQ·차트)"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference, Series
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint as DP
from PIL import Image as PILImage
import importlib
C = importlib.import_module(__import__('os').environ.get('MANUAL_MODULE','content'))

HERE = os.path.dirname(os.path.abspath(__file__))
OPT = os.path.join(HERE, "assets", "opt")
CHAR = os.path.join(HERE, "assets", "mango_char.png")

def solid(color): return PatternFill("solid", fgColor=color.lstrip("#"))
FONT = getattr(C, "FONT", "맑은 고딕")        # 중국어판(content_zh)은 'Microsoft YaHei' 를 쓴다
def F(size=11, bold=False, color="1E2340"): return Font(name=FONT, size=size, bold=bold, color=color.lstrip("#"))
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOPL = Alignment(horizontal="left", vertical="top", wrap_text=True)
thin = Side(style="thin", color="D8DEE9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def clean(s): return str(s).replace("**", "")


def step_key(text, limit=13):
    """스텝 문장에서 플로우 칩용 핵심어 추출 (**볼드** 우선)."""
    t = str(text).strip()
    if "**" in t:
        parts = t.split("**")
        if len(parts) >= 3 and parts[1].strip():
            k = parts[1].strip()
            return k if len(k) <= limit else k[:limit].rstrip() + "…"
    plain = t.replace("**", "")
    for cut in ("—", ".", "!", "?"):
        if cut in plain:
            plain = plain.split(cut)[0]
    plain = plain.strip()
    return plain if len(plain) <= limit else plain[:limit].rstrip() + "…"


def shot_for(sec):
    """섹션 대표 캡처 경로 (imgs 면 첫 장)."""
    key = ""
    if sec.get("imgs"):
        key = sec["imgs"][0][0]
    elif sec.get("img"):
        key = sec["img"]
    if not key:
        return ""
    fp = os.path.join(OPT, key + ".jpg")
    return fp if os.path.exists(fp) else ""

def style_cell(c, fill=None, font=None, align=CEN, border=True):
    if fill: c.fill = fill
    if font: c.font = font
    c.alignment = align
    if border: c.border = BORDER

def build(aud, path):
    t = C.THEMES[aud]; m = C.MANUALS[aud]
    p = t["primary"]; p2 = t["primary2"]; acc = t["accent"]; soft = t["soft"]
    title, subtitle = C.COVERS[aud]
    wb = Workbook()

    # ───── 표지 ─────
    ws = wb.active; ws.title = C.UI["sheet_cover"]; ws.sheet_view.showGridLines = False
    for col in range(1, 12): ws.column_dimensions[get_column_letter(col)].width = 11
    # 상단 배너
    ws.merge_cells("A1:K3")
    ws["A1"].fill = solid(p);
    for r in range(1,4):
        for cc in range(1,12): ws.cell(r,cc).fill = solid(p)
    ws["A1"].value = f"{t['emoji']}  {C.BRAND['name']} · MANGOI"; ws["A1"].font = F(16, True, "FFFFFF"); ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("A5:K6"); ws["A5"].value = title; ws["A5"].font = F(34, True, p); ws["A5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("A7:K7"); ws["A7"].value = subtitle; ws["A7"].font = F(15, True, "475569"); ws["A7"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("A9:F9"); ws["A9"].value = f"{C.UI['target']} · {t['who']}"; ws["A9"].font = F(12, False, "334155"); ws["A9"].alignment = Alignment(horizontal="left", indent=1)
    ws.merge_cells("A10:F10"); ws["A10"].value = f"🌐 {C.SITE}    📞 1644-0561    ✉ support@mangoi.co.kr"; ws["A10"].font = F(11, False, "64748B"); ws["A10"].alignment = Alignment(horizontal="left", indent=1)
    # 안내 박스
    ws.merge_cells("A12:F19"); ws["A12"].fill = solid(soft)
    for r in range(12,20):
        for cc in range(1,7): ws.cell(r,cc).fill = solid(soft)
    ws["A12"].value = C.UI["xls_docuse"]
    ws["A12"].font = F(12, False, "334155"); ws["A12"].alignment = TOPL
    ws.row_dimensions[5].height = 44; ws.row_dimensions[1].height = 22
    # 캐릭터 이미지
    if os.path.exists(CHAR):
        im = PILImage.open(CHAR); ar = im.size[0]/im.size[1]
        xi = XLImage(CHAR); xi.height = 300; xi.width = int(300*ar)
        ws.add_image(xi, "H12")

    # ───── 기능 색인 ─────
    ws2 = wb.create_sheet(C.UI["sheet_index"]); ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:E1"); ws2["A1"].value = C.UI["xls_index_title"].format(label=t["label"]); ws2["A1"].font = F(18, True, p); ws2["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.row_dimensions[1].height = 34
    heads = C.UI["xls_index_head"]
    widths = [7, 8, 22, 62, 40]
    for j, (hd, wd) in enumerate(zip(heads, widths), 1):
        c = ws2.cell(3, j, hd); style_cell(c, solid(p), F(12, True, "FFFFFF"))
        ws2.column_dimensions[get_column_letter(j)].width = wd
    for i, sec in enumerate(m["sections"]):
        r = 4+i
        fill = solid("FFFFFF" if i % 2 == 0 else "F5F7FB")
        style_cell(ws2.cell(r,1, f"{sec['no']:02d}"), fill, F(12, True, p))
        style_cell(ws2.cell(r,2, sec["icon"]), fill, F(16))
        style_cell(ws2.cell(r,3, sec["title"]), fill, F(12, True, "1E2340"), LEFT)
        steps = " → ".join(clean(x) for x in sec["steps"])
        style_cell(ws2.cell(r,4, steps), fill, F(11, False, "334155"), TOPL)
        style_cell(ws2.cell(r,5, clean(sec.get("tip",""))), fill, F(10.5, False, C.BRAND["mango_dark"]), TOPL)
        ws2.row_dimensions[r].height = 54
    ws2.freeze_panes = "A4"

    # ───── 단계별 사용법 ─────
    ws3 = wb.create_sheet(C.UI["sheet_steps"]); ws3.sheet_view.showGridLines = False
    ws3.merge_cells("A1:F1"); ws3["A1"].value = C.UI["xls_steps_title"]; ws3["A1"].font = F(18, True, p); ws3["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws3.row_dimensions[1].height = 34
    for col, wd in zip("ABCDEF", [6, 26, 62, 10, 10, 10]):
        ws3.column_dimensions[col].width = wd
    ws3.column_dimensions["G"].width = 3
    ws3.column_dimensions["H"].width = 58          # 화면 캡처 자리
    r = 3
    for sec in m["sections"]:
        sec_top = r                                 # 이 기능 블록의 첫 행 (캡처 앵커)
        ws3.merge_cells(f"A{r}:F{r}")
        c = ws3.cell(r,1, f"{sec['no']:02d}  {sec['icon']} {sec['title']}  —  {sec['sub']}")
        for cc in range(1,7): ws3.cell(r,cc).fill = solid(p2)
        c.font = F(13, True, "FFFFFF"); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws3.row_dimensions[r].height = 26
        r += 1
        # 설명 문단
        if sec.get("desc"):
            ws3.merge_cells(f"A{r}:F{r}")
            c = ws3.cell(r,1, "📖 " + clean(sec["desc"]))
            c.font = F(11, False, "475569"); c.alignment = TOPL; c.border = BORDER
            for cc in range(1,7): ws3.cell(r,cc).fill = solid("F8FAFC")
            ws3.row_dimensions[r].height = 42
            r += 1
        # 단계 요약(플로우) — 한 줄로 흐름을 먼저 보여 준다
        if sec.get("steps"):
            ws3.merge_cells(f"A{r}:F{r}")
            chain = "   ›   ".join("%d %s" % (i+1, step_key(x)) for i, x in enumerate(sec["steps"]))
            c = ws3.cell(r,1, "🧭  " + chain)
            for cc in range(1,7): ws3.cell(r,cc).fill = solid("EEF2FF")
            c.font = F(10.5, True, p); c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
            c.border = BORDER
            ws3.row_dimensions[r].height = 24
            r += 1
        # 단계 헤더
        style_cell(ws3.cell(r,1,C.UI["xls_step"]), solid(soft), F(11, True, p));
        ws3.merge_cells(f"B{r}:F{r}"); style_cell(ws3.cell(r,2,C.UI["xls_do"]), solid(soft), F(11, True, p), LEFT)
        r += 1
        for k, step in enumerate(sec["steps"], 1):
            style_cell(ws3.cell(r,1, k), None, F(11, True, p2))
            ws3.merge_cells(f"B{r}:F{r}"); style_cell(ws3.cell(r,2, clean(step)), None, F(11, False, "334155"), TOPL)
            ws3.row_dimensions[r].height = 30
            r += 1
        # 핵심 포인트
        for pt in sec.get("points", []):
            style_cell(ws3.cell(r,1,"✅"), None, F(11, False, p))
            ws3.merge_cells(f"B{r}:F{r}"); style_cell(ws3.cell(r,2, clean(pt)), None, F(10.5, False, "1E3A5F"), TOPL)
            for cc in range(1,7): ws3.cell(r,cc).fill = solid("F4F7FF")
            ws3.row_dimensions[r].height = 24
            r += 1
        # 알아두기
        if sec.get("note"):
            ws3.merge_cells(f"A{r}:F{r}")
            c = ws3.cell(r,1, C.UI["note"] + " · " + clean(sec["note"]))
            for cc in range(1,7): ws3.cell(r,cc).fill = solid("FFFBEB")
            c.font = F(10.5, False, "92670A"); c.alignment = TOPL; c.border = BORDER
            ws3.row_dimensions[r].height = 26
            r += 1
        if sec.get("tip"):
            ws3.merge_cells(f"A{r}:F{r}")
            c = ws3.cell(r,1, C.UI["tip"] + " · " + clean(sec["tip"]))
            for cc in range(1,7): ws3.cell(r,cc).fill = solid("FFF7E6")
            c.font = F(11, False, "5B4A2A"); c.alignment = TOPL; c.border = BORDER
            ws3.row_dimensions[r].height = 30
            r += 1
        # 이 기능의 실제 화면 캡처 (H열에 띄운다 — 본문 표를 가리지 않는다)
        fp = shot_for(sec)
        if fp:
            try:
                im = PILImage.open(fp); ar = im.size[0] / im.size[1]
                xi = XLImage(fp); xi.width = 420; xi.height = int(420 / ar)
                ws3.add_image(xi, "H%d" % sec_top)
            except Exception:
                pass
        # 블록 사이 여백 — 캡처(약 263px)가 겹치지 않도록 최소 높이를 확보한다
        used_px = sum(ws3.row_dimensions[x].height or 15 for x in range(sec_top, r))
        pad = 0
        while used_px < 290:
            ws3.row_dimensions[r + pad].height = 18
            used_px += 18; pad += 1
        r += pad
        r += 1  # 여백

    # ───── 체크리스트 ─────
    ws4 = wb.create_sheet(C.UI["sheet_check"]); ws4.sheet_view.showGridLines = False
    cl = m["checklist"]
    ws4.merge_cells("A1:D1"); ws4["A1"].value = C.UI["xls_check_title"].format(title=cl["title"]); ws4["A1"].font = F(18, True, p); ws4["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws4.row_dimensions[1].height = 34
    ws4.column_dimensions["A"].width = 6; ws4.column_dimensions["B"].width = 70; ws4.column_dimensions["C"].width = 14
    style_cell(ws4.cell(3,1,C.UI["xls_check_head"][0]), solid(p), F(12,True,"FFFFFF"))
    style_cell(ws4.cell(3,2,C.UI["xls_check_head"][1]), solid(p), F(12,True,"FFFFFF"), LEFT)
    style_cell(ws4.cell(3,3,C.UI["xls_check_head"][2]), solid(p), F(12,True,"FFFFFF"))
    for i, it in enumerate(cl["items"]):
        r = 4+i; fill = solid("FFFFFF" if i%2==0 else "F5F7FB")
        style_cell(ws4.cell(r,1,"☐"), fill, F(14, False, p))
        style_cell(ws4.cell(r,2, clean(it)), fill, F(12, False, "334155"), LEFT)
        style_cell(ws4.cell(r,3,""), fill)
        ws4.row_dimensions[r].height = 30

    # ───── FAQ ─────
    ws5 = wb.create_sheet(C.UI["sheet_faq"]); ws5.sheet_view.showGridLines = False
    ws5.merge_cells("A1:B1"); ws5["A1"].value = C.UI["xls_faq_title"]; ws5["A1"].font = F(18, True, p); ws5["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[1].height = 34
    ws5.column_dimensions["A"].width = 34; ws5.column_dimensions["B"].width = 74
    style_cell(ws5.cell(3,1,C.UI["xls_faq_head"][0]), solid(p), F(12,True,"FFFFFF"), LEFT)
    style_cell(ws5.cell(3,2,C.UI["xls_faq_head"][1]), solid(p), F(12,True,"FFFFFF"), LEFT)
    for i,(q,a) in enumerate(m["faq"]):
        r=4+i; fill = solid("FFFFFF" if i%2==0 else "F5F7FB")
        style_cell(ws5.cell(r,1, clean(q)), fill, F(12, True, p2), TOPL)
        style_cell(ws5.cell(r,2, clean(a)), fill, F(11, False, "334155"), TOPL)
        ws5.row_dimensions[r].height = 46

    # ───── 데이터·차트 ─────
    ws6 = wb.create_sheet(C.UI["sheet_data"]); ws6.sheet_view.showGridLines = False
    ws6.merge_cells("A1:D1"); ws6["A1"].value = C.UI["xls_data_title"]; ws6["A1"].font = F(18, True, p); ws6["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws6.row_dimensions[1].height = 34
    ws6.column_dimensions["A"].width = 22; ws6.column_dimensions["B"].width = 12
    charts = list(m["charts"].values())
    # 표1 (bar/column)
    c1 = charts[1]
    ws6.cell(3,1, c1["title"]).font = F(12, True, p)
    style_cell(ws6.cell(4,1,C.UI["xls_item"]), solid(p), F(11,True,"FFFFFF")); style_cell(ws6.cell(4,2,C.UI["xls_value"]), solid(p), F(11,True,"FFFFFF"))
    for i,(lab,val) in enumerate(zip(c1["labels"], c1["values"])):
        style_cell(ws6.cell(5+i,1, lab), None, F(11,False,"334155"), LEFT)
        style_cell(ws6.cell(5+i,2, val), None, F(11,True,p))
    n1 = len(c1["labels"])
    bar = BarChart(); bar.type="col"; bar.title=c1["title"]; bar.height=8; bar.width=15
    data = Reference(ws6, min_col=2, min_row=4, max_row=4+n1)
    cats = Reference(ws6, min_col=1, min_row=5, max_row=4+n1)
    bar.add_data(data, titles_from_data=True); bar.set_categories(cats)
    bar.legend = None; bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True
    s0 = bar.series[0]; s0.graphicalProperties.solidFill = c1["colors"][0].lstrip("#")
    ws6.add_chart(bar, "D3")

    # 표2 (doughnut) — 있으면
    if len(charts) > 2:
        c2 = charts[2]; base = 5+n1+2
        ws6.cell(base,1, c2["title"]).font = F(12, True, p)
        style_cell(ws6.cell(base+1,1,C.UI["xls_item"]), solid(p2), F(11,True,"FFFFFF")); style_cell(ws6.cell(base+1,2,C.UI["xls_value"]), solid(p2), F(11,True,"FFFFFF"))
        for i,(lab,val) in enumerate(zip(c2["labels"], c2["values"])):
            style_cell(ws6.cell(base+2+i,1, lab), None, F(11,False,"334155"), LEFT)
            style_cell(ws6.cell(base+2+i,2, val), None, F(11,True,p2))
        n2 = len(c2["labels"])
        dough = DoughnutChart(); dough.title=c2["title"]; dough.height=8; dough.width=15
        d2 = Reference(ws6, min_col=2, min_row=base+1, max_row=base+1+n2)
        ct2 = Reference(ws6, min_col=1, min_row=base+2, max_row=base+1+n2)
        dough.add_data(d2, titles_from_data=True); dough.set_categories(ct2)
        dough.dataLabels = DataLabelList(); dough.dataLabels.showVal = True
        ser = dough.series[0]
        for i, colr in enumerate(c2["colors"][:n2]):
            pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = colr.lstrip("#")
            ser.data_points.append(pt)
        ws6.add_chart(dough, f"D{base}")

    # ───── 마무리(결론) — 있으면 ─────
    if m.get("conclusion"):
        cc = m["conclusion"]
        ws7 = wb.create_sheet(C.UI["concl_page"]); ws7.sheet_view.showGridLines = False
        ws7.column_dimensions["A"].width = 22; ws7.column_dimensions["B"].width = 86
        ws7.merge_cells("A1:B1"); ws7["A1"].value = "🏁 " + cc["title"]; ws7["A1"].font = F(17, True, p)
        ws7["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1); ws7.row_dimensions[1].height = 32
        ws7.merge_cells("A2:B2"); ws7["A2"].value = clean(cc["lead"]); ws7["A2"].font = F(11, False, "334155"); ws7["A2"].alignment = TOPL; ws7.row_dimensions[2].height = 54
        r = 4
        ws7.merge_cells(f"A{r}:B{r}"); c = ws7.cell(r,1, C.UI["concl_flow"]); c.font = F(13, True, "FFFFFF")
        for cc_ in (1,2): ws7.cell(r,cc_).fill = solid(p2)
        c.alignment = Alignment(horizontal="left", indent=1); ws7.row_dimensions[r].height = 24; r += 1
        for stg, act in cc["flow"]:
            style_cell(ws7.cell(r,1, stg), solid(soft), F(11, True, p), LEFT)
            style_cell(ws7.cell(r,2, clean(act)), None, F(11, False, "334155"), TOPL); ws7.row_dimensions[r].height = 30; r += 1
        r += 1
        ws7.merge_cells(f"A{r}:B{r}"); c = ws7.cell(r,1, C.UI["concl_principles"]); c.font = F(13, True, "FFFFFF")
        for cc_ in (1,2): ws7.cell(r,cc_).fill = solid(p)
        c.alignment = Alignment(horizontal="left", indent=1); ws7.row_dimensions[r].height = 24; r += 1
        for i, x in enumerate(cc["principles"], 1):
            style_cell(ws7.cell(r,1, i), None, F(11, True, p2)); style_cell(ws7.cell(r,2, clean(x)), None, F(11, False, "334155"), TOPL); ws7.row_dimensions[r].height = 26; r += 1
        r += 1
        ws7.merge_cells(f"A{r}:B{r}"); c = ws7.cell(r,1, C.UI["concl_remember"]); c.font = F(13, True, "FFFFFF")
        for cc_ in (1,2): ws7.cell(r,cc_).fill = solid(C.BRAND["mango_dark"])
        c.alignment = Alignment(horizontal="left", indent=1); ws7.row_dimensions[r].height = 24; r += 1
        for x in cc["remember"]:
            style_cell(ws7.cell(r,1, "📌"), solid("FFFBEB"), F(11)); style_cell(ws7.cell(r,2, clean(x)), solid("FFFBEB"), F(11, False, "92670A"), TOPL); ws7.row_dimensions[r].height = 26; r += 1
        r += 1
        ws7.merge_cells(f"A{r}:B{r}"); c = ws7.cell(r,1, "🍊 " + clean(cc["closing"]))
        for cc_ in (1,2): ws7.cell(r,cc_).fill = solid(soft)
        c.font = F(11.5, True, p); c.alignment = TOPL; c.border = BORDER; ws7.row_dimensions[r].height = 66

    wb.save(path)
    print("XLSX", aud, "->", path)

if __name__ == "__main__":
    suffix = getattr(C, "OUT_SUFFIX", "")
    outdir = os.path.join(HERE, "build_xlsx" + suffix); os.makedirs(outdir, exist_ok=True)
    for aud in list(C.MANUALS.keys()):
        build(aud, os.path.join(outdir, C.FILENAMES[aud] + ".xlsx"))
